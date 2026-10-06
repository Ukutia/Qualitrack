import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
import docx, openpyxl
st=getSampleStyleSheet()
filler="La universidad mantiene registros administrativos de matrícula, horarios de salas y uso de laboratorios durante el semestre. "*6
key=("Resolución Exenta N° 123: se aprueba la Política de Aseguramiento Interno de la Calidad de la Universidad, "
     "coherente con la misión y los propósitos institucionales. Se crea la Dirección de Calidad como responsable de su implementación, "
     "según consta en el organigrama institucional y en el Plan de Desarrollo Estratégico (PDE) 2025-2030.")
# PDF 3 pages, key content on page 3
d=SimpleDocTemplate("prueba_politica.pdf",pagesize=letter)
el=[Paragraph("Informe institucional 2026",st['Title'])]
for p in range(2):
    el+= [Paragraph(filler,st['Normal']) for _ in range(6)]+[PageBreak()]
el+=[Paragraph(filler,st['Normal']),Paragraph(key,st['Normal']),Paragraph(filler,st['Normal'])]
d.build(el)
# DOCX
doc=docx.Document(); doc.add_heading("Acta de consejo académico",1)
for _ in range(15): doc.add_paragraph(filler)
doc.add_page_break(); doc.add_paragraph(key)
for _ in range(5): doc.add_paragraph(filler)
doc.save("prueba_acta.docx")
# XLSX
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Resumen"
ws.append(["Año","Matrícula","Sede"])
for i in range(40): ws.append([2000+i, 1000+i*7, "Central"])
ws2=wb.create_sheet("Calidad")
for i in range(12): ws2.append([f"Ítem {i}", "Registro de uso de laboratorios"])
ws2["C9"]=key
wb.save("prueba_indicadores.xlsx")
# >10MB pdf
with open("grande_11mb.pdf","wb") as f: f.write(open("prueba_politica.pdf","rb").read()+os.urandom(11*1024*1024))
